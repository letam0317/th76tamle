"""
sheet-bai-xe.py — g-sheet "BÃI XE NCC" (luật BX-01): mỗi lượt xe vào bãi bảo vệ/NCC = 1 dòng kèm ảnh, bạn quyết ở cột VÀNG.

  User 02/10/2026: "sao k đưa vô g-sheet > t quyết định trên đó và bạn canh quyết định từ tôi để chạy task > chuyển sang b1.1 luôn".
  Nguồn: Videos/5S-bang-chung/bai-xe/<ngày>/ket-qua.json (canh-bai-xe.mjs) — 14 ngày gần nhất, mới nhất lên trước.
  Quyết định của bạn đọc lại theo "Mã lượt" ⇒ không mất qua các lần dựng lại. Trạng thái máy làm lấy từ .exports/bai-xe-phieu.json.
  Sheet riêng của letam0317, KHÔNG chia sẻ (ảnh có mặt người + ảnh cổng bảo vệ). Ảnh nhúng thẳng vào file, không cần link công khai.

  Chạy:  python sheet-bai-xe.py          (diễn tập: dựng .xlsx tạm, in tình trạng)
         python sheet-bai-xe.py --doc    (chỉ đọc quyết định trên Sheet → .exports/bai-xe-quyet-dinh.json, không nạp đè)
         python sheet-bai-xe.py --lam    (đọc quyết định + dựng lại + nạp đè + đọc lại kiểm)
"""
import io, json, os, sys, glob, datetime as dt, urllib.request, urllib.parse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")
DIR = os.path.dirname(os.path.abspath(__file__))
EXP = os.path.join(DIR, ".exports")
GOC = os.environ.get("BAIXE_DIR", r"C:\Users\lechitam\Videos\5S-bang-chung\bai-xe")
SO_ID = os.path.join(EXP, "bai-xe-sheet.json")
SO_PHIEU = os.path.join(EXP, "bai-xe-phieu.json")
RA_QD = os.path.join(EXP, "bai-xe-quyet-dinh.json")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TAB = "BÃI XE NCC"
LAM, DOC = "--lam" in sys.argv, "--doc" in sys.argv
now = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None) + dt.timedelta(hours=7)
QD = ["Lập phiếu (NV Hasaki)", "Không phải NV Hasaki", "Xe ra / không vi phạm", "Không rõ – bỏ qua"]

def token():
    rc = json.load(open(os.path.join(os.path.expanduser("~"), ".clasprc.json"), encoding="utf8"))
    t = rc["tokens"]["default"] if "tokens" in rc else rc["token"]
    body = urllib.parse.urlencode({"client_id": t["client_id"], "client_secret": t["client_secret"],
                                   "refresh_token": t["refresh_token"], "grant_type": "refresh_token"}).encode()
    return json.load(urllib.request.urlopen("https://oauth2.googleapis.com/token", body))["access_token"]
TK = None
def http(url, data=None, method=None, headers=None):
    h = {"authorization": "Bearer " + TK, **(headers or {})}
    return urllib.request.urlopen(urllib.request.Request(url, data=data, method=method, headers=h)).read()
def multipart(meta, buf):
    b = "baixe" + now.strftime("%H%M%S")
    body = (b"--" + b.encode() + b"\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n" + json.dumps(meta).encode() +
            b"\r\n--" + b.encode() + b"\r\nContent-Type: " + XLSX.encode() + b"\r\n\r\n" + buf + b"\r\n--" + b.encode() + b"--")
    return body, {"content-type": "multipart/related; boundary=" + b}

doc_json = lambda f, d: json.load(open(f, encoding="utf8")) if os.path.exists(f) else d
sheet_id = doc_json(SO_ID, {}).get("id")

# ── đọc quyết định đang có trên Sheet ──────────────────────────────────────────────────────────────
qd_cu = {}
if (LAM or DOC) and sheet_id:
    TK = token()
    wb0 = load_workbook(io.BytesIO(http("https://www.googleapis.com/drive/v3/files/%s/export?mimeType=%s" % (sheet_id, urllib.parse.quote(XLSX)))))
    ws0 = wb0[TAB] if TAB in wb0.sheetnames else wb0.worksheets[0]
    dau = next((r for r in range(1, 15) if ws0.cell(r, 1).value == "Mã lượt"), None)
    if dau:
        ten = [ws0.cell(dau, c).value for c in range(1, ws0.max_column + 1)]
        for r in range(dau + 1, ws0.max_row + 1):
            v = {ten[c - 1]: ws0.cell(r, c).value for c in range(1, len(ten) + 1) if ten[c - 1]}
            if v.get("Mã lượt"):
                qd_cu[str(v["Mã lượt"])] = {"qd": v.get("Bạn quyết") or "", "ma": str(v.get("Mã NV vi phạm") or "").strip().replace(".0", ""),
                                            "gc": v.get("Ghi chú của bạn") or ""}
json.dump(qd_cu, open(RA_QD, "w", encoding="utf8"), ensure_ascii=False, indent=1)
# quyết định bấm trên dashboard bang-591 → hiện cả trên sheet (dashboard ghi sau thì thắng)
for k, w in doc_json(os.path.join(EXP, "bang-591-quyet-dinh.json"), {}).items():
    qd_cu[k] = {"qd": w.get("qd", ""), "ma": w.get("ma", ""), "gc": w.get("gc", "") or qd_cu.get(k, {}).get("gc", "")}
if DOC:
    n = sum(1 for x in qd_cu.values() if x["qd"])
    print("Đọc %d dòng · %d quyết định → %s" % (len(qd_cu), n, RA_QD)); sys.exit(0)

# ── ảnh gốc lên Drive RIÊNG (không chia sẻ) để mở to: 1 lần/ảnh, nhớ id ở .exports/bai-xe-anh-drive.json ──
SO_ANH = os.path.join(EXP, "bai-xe-anh-drive.json"); anh_id = doc_json(SO_ANH, {})
def anh_drive(f, ma):
    global TK
    if not LAM or not os.path.exists(f): return anh_id.get(ma, {}).get("id")
    m = anh_id.get(ma)
    if m and m.get("mtime") == os.path.getmtime(f): return m["id"]
    TK = TK or token()
    so = doc_json(SO_ID, {})
    if not so.get("thuMucAnh"):
        so["thuMucAnh"] = json.loads(http("https://www.googleapis.com/drive/v3/files", json.dumps({"name": "BAI-XE-NCC-anh (5S, riêng tư)",
            "mimeType": "application/vnd.google-apps.folder"}).encode(), "POST", {"content-type": "application/json"}))["id"]
        json.dump(so, open(SO_ID, "w", encoding="utf8"))
    bd = "anh" + now.strftime("%H%M%S")
    meta = {"name": ma + ".jpg", "parents": [so["thuMucAnh"]]} if not m else {}
    body = (b"--" + bd.encode() + b"\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n" + json.dumps(meta).encode() +
            b"\r\n--" + bd.encode() + b"\r\nContent-Type: image/jpeg\r\n\r\n" + open(f, "rb").read() + b"\r\n--" + bd.encode() + b"--")
    h = {"content-type": "multipart/related; boundary=" + bd}
    fid = json.loads(http("https://www.googleapis.com/upload/drive/v3/files/%s?uploadType=multipart" % m["id"], body, "PATCH", h))["id"] if m else           json.loads(http("https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart", body, "POST", h))["id"]
    anh_id[ma] = {"id": fid, "mtime": os.path.getmtime(f)}; json.dump(anh_id, open(SO_ANH, "w", encoding="utf8"), indent=1)
    return fid

# ── dựng bảng ──────────────────────────────────────────────────────────────────────────────────────
phieu = doc_json(SO_PHIEU, {})
luot = []
for f in sorted(glob.glob(os.path.join(GOC, "*", "ket-qua.json")), reverse=True)[:14]:
    k = doc_json(f, {}); ngay = k.get("ngay")
    for l in k.get("luot", []):
        if l.get("id"): luot.append((ngay, os.path.dirname(f), l))
luot.sort(key=lambda x: (x[0], x[2].get("den", "")), reverse=True)

wb = Workbook(); ws = wb.active; ws.title = TAB
F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
TOP = Alignment(vertical="top", wrap_text=True); BD = Border(*(Side(style="thin", color="D0D4DA"),) * 4)
VANG = "FFF2B3"
cho = [l for (_, _, l) in luot if not qd_cu.get(l["id"], {}).get("qd")]
lam_xong = [i for i, p in phieu.items() if p.get("trangThai") == "xong"]
ws["A1"] = "BÃI XE NCC / BẢO VỆ — lượt NV có thể gửi xe sai bãi (luật BX-01 · mục 25 · trừ 2%)"; ws["A1"].font = F(bold=True, size=14)
ws["A2"] = ("%d lượt (14 ngày) · %d chờ bạn quyết · %d đã lập phiếu · cập nhật %s" % (len(luot), len(cho), len(lam_xong), now.strftime("%H:%M %d/%m")))
ws["A3"] = ("Cách dùng: xem ảnh (khung đỏ = bãi NCC/bảo vệ; hàng dưới = ảnh cổng người vào ngay sau — CHỈ để đối chiếu, không lên phiếu) → "
            "chọn ở cột VÀNG 'Bạn quyết'. Chọn «%s» thì gõ MÃ NV ở cột 'Mã NV vi phạm' → máy (30 phút/lần) tự dựng clip khung bám người, "
            "lập phiếu 5S và chuyển B1.1 cho người đó; kết quả + link phiếu hiện ở cột 'Máy đã làm'." % QD[0])
for c in ("A2", "A3"): ws[c].alignment = TOP; ws[c].font = F(color="444444")
ws.merge_cells("A3:L3"); ws.row_dimensions[3].height = 46
R0 = 5
COT = [("Mã lượt", 22), ("Ngày", 11), ("Camera", 12), ("Giờ thấy người", 10), ("Ảnh lượt (bấm cột F để xem ảnh to)", 128), ("Ảnh to", 12), ("Chấm công gần giờ", 30),
       ("Sổ cổng gần giờ", 28), ("Bạn quyết", 22), ("Mã NV vi phạm", 13), ("Ghi chú của bạn", 24), ("Máy đã làm", 34)]
for j, (h, w) in enumerate(COT, 1):
    c = ws.cell(R0, j, h); c.font = F(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79"); c.alignment = TOP; c.border = BD
    ws.column_dimensions[chr(64 + j)].width = w
for j in (9, 10, 11): ws.cell(R0, j).fill = PatternFill("solid", fgColor="B8860B")
tam = os.path.join(EXP, ".bai-xe-tam", "_anh-sheet"); os.makedirs(tam, exist_ok=True)
for n, (ngay, thu, l) in enumerate(luot):
    r = R0 + 1 + n; cu = qd_cu.get(l["id"], {}); p = phieu.get(l["id"], {})
    ds = lambda a: "\n".join("%s · %s %s" % (x["gio"][:5], x["code"], x["ten"]) for x in a) or "—"
    may = "—"
    if p.get("trangThai") == "xong": may = "✓ phiếu %s → B1.1 %s\nhttps://work.hasaki.vn/tasks?task_id=%s" % (p.get("phieu"), p.get("nv", ""), p.get("phieu"))
    elif p.get("trangThai"): may = "%s: %s" % (p["trangThai"], p.get("loi", ""))
    fid = anh_drive(os.path.join(thu, l["anh"]), l["id"])
    to = '=HYPERLINK("https://drive.google.com/file/d/%s/view","🔍 Mở ảnh to")' % fid if fid else "—"
    val = [l["id"], ngay, l["camera"], l.get("gioNguoi") or "—", "", to, ds(l.get("chamCong", [])), ds(l.get("cong", [])),
           cu.get("qd", ""), cu.get("ma", ""), cu.get("gc", ""), may]
    for j, x in enumerate(val, 1):
        c = ws.cell(r, j, x); c.alignment = TOP; c.border = BD; c.font = F(bold=(j == 1))
    for j in (9, 10, 11): ws.cell(r, j).fill = PatternFill("solid", fgColor=VANG)
    if fid: ws.cell(r, 6).font = F(bold=True, color="1155CC", underline="single")
    if p.get("trangThai") == "xong":
        for j in range(1, 13):
            if j not in (6, 9, 10, 11): ws.cell(r, j).font = F(color="888888")
    anh = os.path.join(thu, l["anh"])
    if os.path.exists(anh):
        im = Image.open(anh); w = 900; h = int(im.height * w / im.width)
        nho = os.path.join(tam, l["id"] + ".jpg"); im.convert("RGB").resize((w, h)).save(nho, quality=70)
        xi = XLImage(nho); xi.anchor = "E%d" % r; ws.add_image(xi)
        ws.row_dimensions[r].height = h * 0.75 + 6
dv = DataValidation(type="list", formula1='"%s"' % ",".join(QD), allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv); dv.add("I%d:I%d" % (R0 + 1, R0 + len(luot) + 50))
ws.freeze_panes = "B%d" % (R0 + 1)
buf = io.BytesIO(); wb.save(buf); buf = buf.getvalue()
open(os.path.join(EXP, ".bai-xe-tam", "sheet-bai-xe.xlsx"), "wb").write(buf)
print("Dựng: %d lượt · %d chờ quyết · %d đã lập phiếu · %.1f MB" % (len(luot), len(cho), len(lam_xong), len(buf) / 1048576))
if not LAM:
    print("(diễn tập — thêm --lam để nạp lên Sheet)"); sys.exit(0)

TK = TK or token()
if not sheet_id:
    body, h = multipart({"name": "BÃI XE NCC – quyết định (5S)", "mimeType": "application/vnd.google-apps.spreadsheet"}, buf)
    sheet_id = json.loads(http("https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart", body, "POST", h))["id"]
    json.dump({"id": sheet_id, "tao": now.strftime("%Y-%m-%d %H:%M")}, open(SO_ID, "w", encoding="utf8"))
    print("Tạo Sheet mới (riêng tư, không chia sẻ):", sheet_id)
else:
    body, h = multipart({"mimeType": "application/vnd.google-apps.spreadsheet"}, buf)
    http("https://www.googleapis.com/upload/drive/v3/files/%s?uploadType=multipart" % sheet_id, body, "PATCH", h)
# đọc lại kiểm
wb1 = load_workbook(io.BytesIO(http("https://www.googleapis.com/drive/v3/files/%s/export?mimeType=%s" % (sheet_id, urllib.parse.quote(XLSX)))))
w1 = wb1.worksheets[0]
dau = next((r for r in range(1, 15) if w1.cell(r, 1).value == "Mã lượt"), None)
so = sum(1 for r in range(dau + 1, w1.max_row + 1) if w1.cell(r, 1).value) if dau else 0
giu = sum(1 for r in range(dau + 1, w1.max_row + 1) if dau and w1.cell(r, 9).value)
print("Đọc lại: %d dòng · %d quyết định · ảnh %d → %s" % (so, giu, len(getattr(w1, "_images", [])), "OK" if so == len(luot) else "LỆCH"))
print("Link: https://docs.google.com/spreadsheets/d/%s/edit" % sheet_id)
